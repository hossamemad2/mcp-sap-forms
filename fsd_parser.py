"""
parse_fsd: extracts candidate interface fields from a Functional Spec
Document (.docx / .xlsx / .pdf).

Classifies each extracted field into a `section` -- "HEADER", "ITEM", or
None (flat/top-level, for simple forms like the original ZTRM_LG_F01
letter that have no repeating structure) -- by keyword-matching the
heading/sheet-name text nearest to the table it came from ("Header",
"General Data" -> HEADER; "Item", "Line Item", "Detail", "Position" ->
ITEM). This is a heuristic, not a guarantee: anything the heuristic can't
confidently classify is still returned (never silently dropped), tagged
section=None and flagged in `needs_review`, so the calling agent can
resolve ambiguous cases by reading the surrounding document text rather
than the tool guessing wrong silently.
"""
from __future__ import annotations

import os
import re

HEADER_KEYWORDS = ("header", "general data", "general information", "main data", "top level")
ITEM_KEYWORDS = ("item", "line item", "detail", "position", "row", "schedule line")

FIELD_NAME_COLS = ("field name", "field", "name", "parameter", "technical name")
TYPE_COLS = ("type", "data type", "abap type", "datatype")
DESC_COLS = ("description", "desc", "label", "text")


def classify_section(heading_text: str):
    if not heading_text:
        return None
    t = heading_text.lower()
    if any(k in t for k in ITEM_KEYWORDS):
        return "ITEM"
    if any(k in t for k in HEADER_KEYWORDS):
        return "HEADER"
    return None


def _normalize_field_name(raw: str) -> str:
    name = re.sub(r"[^A-Za-z0-9_]", "_", raw.strip()).upper()
    name = re.sub(r"_+", "_", name).strip("_")
    return name


def _guess_typename(type_hint: str) -> str:
    if not type_hint:
        return "STRING"
    t = type_hint.strip().upper()
    mapping = {
        "STRING": "STRING", "CHAR": "STRING", "TEXT": "STRING",
        "DATE": "DATS", "DATS": "DATS",
        "AMOUNT": "STRING", "CURR": "STRING", "DECIMAL": "STRING", "NUMBER": "STRING", "QTY": "STRING",
        "FLAG": "ABAP_BOOL", "BOOLEAN": "ABAP_BOOL",
    }
    for key, val in mapping.items():
        if key in t:
            return val
    return "STRING"  # safe default; the reviewing agent can override per field


def _row_to_field(headers, row, section):
    cells = {h: (row[i] if i < len(row) else "") for i, h in enumerate(headers)}

    def find(cols):
        for h in headers:
            hl = (h or "").strip().lower()
            if any(c in hl for c in cols):
                return (cells.get(h) or "").strip()
        return ""

    raw_name = find(FIELD_NAME_COLS)
    if not raw_name:
        return None
    return {
        "name": _normalize_field_name(raw_name),
        "source_label": raw_name,
        "type_hint": find(TYPE_COLS),
        "typename_guess": _guess_typename(find(TYPE_COLS)),
        "description": find(DESC_COLS),
        "section": section,
        "needs_review": section is None or not find(TYPE_COLS),
    }


def _looks_like_field_table(headers):
    joined = " ".join((h or "").lower() for h in headers)
    return any(c in joined for c in FIELD_NAME_COLS)


# ---------------------------------------------------------------- DOCX ---

def _parse_docx(path):
    import docx  # python-docx

    document = docx.Document(path)
    body = document.element.body
    current_heading = ""
    fields = []
    raw_blocks = []

    # python-docx tables/paragraphs are separate collections; walk the
    # underlying XML body so we see them in document order and can track
    # "most recent heading paragraph" per table.
    table_index = 0
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            text = "".join(node.text or "" for node in child.iter()
                            if node.tag.endswith("}t"))
            if text.strip():
                current_heading = text.strip()
        elif tag == "tbl":
            table = document.tables[table_index]
            table_index += 1
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            headers = rows[0]
            if _looks_like_field_table(headers):
                section = classify_section(current_heading)
                for row in rows[1:]:
                    f = _row_to_field(headers, row, section)
                    if f:
                        fields.append(f)
            else:
                raw_blocks.append({"heading": current_heading, "rows": rows})

    return fields, raw_blocks


# ---------------------------------------------------------------- XLSX ---

def _parse_xlsx(path):
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    fields = []
    raw_blocks = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = [[("" if c.value is None else str(c.value)) for c in row] for row in ws.iter_rows()]
        rows = [r for r in rows if any(cell.strip() for cell in r)]
        if not rows:
            continue
        headers = rows[0]
        section = classify_section(sheet_name)
        if _looks_like_field_table(headers):
            for row in rows[1:]:
                f = _row_to_field(headers, row, section)
                if f:
                    fields.append(f)
        else:
            raw_blocks.append({"heading": sheet_name, "rows": rows})

    return fields, raw_blocks


# ----------------------------------------------------------------- PDF ---

def _parse_pdf(path):
    import pdfplumber

    fields = []
    raw_blocks = []

    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            words = page.extract_words()
            tables = page.find_tables()
            for t in tables:
                heading_words = [w["text"] for w in words if w["top"] < t.bbox[1] and w["top"] > t.bbox[1] - 40]
                heading = " ".join(heading_words[-8:])  # a short window right above the table
                rows = t.extract()
                if not rows:
                    continue
                headers = rows[0]
                section = classify_section(heading)
                if _looks_like_field_table(headers):
                    for row in rows[1:]:
                        f = _row_to_field(headers, row, section)
                        if f:
                            fields.append(f)
                else:
                    raw_blocks.append({"heading": heading, "rows": rows})

    return fields, raw_blocks


def parse_fsd(path: str) -> dict:
    """
    Returns {"fields": [...], "raw_blocks": [...], "warnings": [...]}.
    `fields` entries always have a `name`; `typename_guess` is a best-effort
    STRING/DATS/ABAP_BOOL guess the calling agent should confirm against
    the FSD's actual type column before using it in deploy_form -- never
    trust typename_guess blindly for anything that isn't obviously text.
    `raw_blocks` holds tables/sections the heuristic couldn't classify as
    a field list at all, returned as-is so nothing is silently dropped.
    """
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        fields, raw_blocks = _parse_docx(path)
    elif ext in (".xlsx", ".xlsm"):
        fields, raw_blocks = _parse_xlsx(path)
    elif ext == ".pdf":
        fields, raw_blocks = _parse_pdf(path)
    else:
        raise ValueError(f"Unsupported FSD format: {ext} (expected .docx, .xlsx, or .pdf)")

    warnings = []
    if not fields:
        warnings.append("No field table detected automatically -- check raw_blocks and "
                         "classify fields manually.")
    unclassified = [f["name"] for f in fields if f["section"] is None]
    if unclassified:
        warnings.append(f"{len(unclassified)} field(s) have no HEADER/ITEM section detected: "
                         f"{', '.join(unclassified[:10])}{'...' if len(unclassified) > 10 else ''} "
                         f"-- confirm whether these are flat top-level fields or belong to a section.")

    return {"fields": fields, "raw_blocks": raw_blocks, "warnings": warnings}
